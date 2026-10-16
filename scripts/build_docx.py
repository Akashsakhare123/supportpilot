#!/usr/bin/env python3
"""Build the SupportPilot project document (Word) from docs/plan/tasks.json + docs/PROGRESS.md.

Called by `python scripts/progress.py sync`; can also run alone:
    python scripts/build_docx.py [output.docx]
Requires python-docx.
"""
from __future__ import annotations

import datetime as dt
import json
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Mm, Cm

BRAND = RGBColor(0x3B, 0x3F, 0x8C)
BRAND_HEX = "3B3F8C"
INK = RGBColor(0x22, 0x24, 0x33)
MUTED = RGBColor(0x5F, 0x63, 0x7A)
STATUS_FILL = {"Done": "D9F2E3", "In progress": "FFF1CC", "Not started": "F1F2F6"}
STATUS_INK = {"Done": RGBColor(0x1E, 0x6B, 0x3A), "In progress": RGBColor(0x8A, 0x5A, 0x00), "Not started": MUTED}
CONTENT_W = 170  # mm (A4 210 - 2 x 20 margins)

ROOT = Path(__file__).resolve().parent.parent

# ------------------------------------------------------------ static text ---
SUMMARY = (
    "SupportPilot is a multi-tenant SaaS platform for AI customer support. A business uploads its "
    "knowledge base, pastes one script tag on its website, and its customers get instant answers grounded "
    "in that business's own documents, with sources. When the AI is not confident, or the customer asks for "
    "a person, the conversation is handed to a human agent in real time."
)
AUDIENCE = [
    ("For clients", "a white-label support widget they can brand with their colors, logo and greeting, "
     "billed as a monthly subscription."),
    ("For hiring companies", "a production-style full-stack project: multi-tenancy, RAG, real-time sockets, "
     "background jobs, billing, security, tests and CI/CD, deployed live."),
]
PROBLEM = [
    "Small and mid-size businesses answer the same questions again and again by email and chat.",
    "Generic chatbots make things up because they are not grounded in the business's own docs.",
    "Pure AI bots frustrate customers when they cannot reach a person.",
]
SOLUTION = [
    "Answers come only from the business's own knowledge base, with source citations.",
    "Every answer carries a confidence score; low confidence escalates to a human automatically.",
    "Agents work from a live inbox and can hand the conversation back to the AI.",
    "Analytics show what the AI resolved and which questions it could not answer, so the client knows which docs to add.",
]
FEATURES = [
    ("Knowledge base", "Upload PDF, text / markdown or a URL; ingested in the background into pgvector."),
    ("Grounded AI answers", "RAG pipeline, streamed token by token, with citations and a confidence score."),
    ("Embeddable widget", "One script tag; tiny Preact bundle in Shadow DOM, themed per tenant."),
    ("Human handoff", "Confidence threshold or 'talk to a human' moves the chat to the live agent inbox."),
    ("Live chat", "Typing indicators, agent presence, message history, reconnect sync."),
    ("Analytics", "AI vs. human resolution, top questions, unanswered questions report."),
    ("Billing", "Stripe subscriptions with plan limits enforced from usage events."),
    ("Security", "Tenant isolation (scoped queries + Postgres RLS), rate limiting, prompt-injection and PII safeguards."),
]
STANDOUT = [
    ("Unanswered questions report", "Shows which questions the AI failed on, so the client knows what docs to add", "T12"),
    ("Prompt-injection and PII safeguards", "Security in the chat pipeline, which most portfolio projects skip", "T15"),
    ("Widget theming per tenant", "Colors, logo, greeting: white-label ready for clients", "T03, T08"),
    ("Eval set", "Test questions that score answer quality in CI", "T14"),
]
CHAT_FLOW = [
    "Customer opens the widget and connects over WebSocket, identified by the tenant's public key.",
    "The message goes to the backend, which embeds it and searches pgvector for the top chunks from that tenant's docs.",
    "The LLM answers using only those chunks and returns a confidence score and sources.",
    "High confidence: the answer is sent. Low confidence, or the customer asks for a human: the conversation is "
    "flagged 'needs agent' and appears in the agent inbox in real time.",
    "The agent replies from the dashboard; the AI stays out of that conversation until it is handed back.",
]
STATES = [
    ("ai", "AI answers from the knowledge base.", "waiting (low confidence / customer asks), closed"),
    ("waiting", "Flagged 'needs agent'; customer sees 'connecting you to an agent'.", "human (agent claims), closed (no agent, email captured)"),
    ("human", "Agent replies; AI pipeline is skipped.", "ai (hand back), closed"),
    ("closed", "Conversation finished.", "-"),
]
COMPONENTS = [
    ("Widget", "Branded chat UI on the client's site; connects with the tenant public key.", "Preact + Vite, Shadow DOM"),
    ("Dashboard", "Signup, onboarding, knowledge base, agent inbox, analytics, settings, billing.", "Next.js on Vercel"),
    ("API", "Auth, tenants, documents, RAG chat, conversations, handoff, analytics, billing webhooks.", "NestJS + Socket.io on Railway"),
    ("Worker", "Extract, chunk, embed and store documents off the request path, with retries.", "BullMQ on Railway"),
    ("PostgreSQL + pgvector", "Relational data and embeddings in one database.", "Railway Postgres"),
    ("Redis", "Queues, Socket.io adapter, presence, rate limits, cache.", "Railway Redis"),
]
TENANCY = [
    ("1. Request context", "Tenant id comes only from the verified JWT or the widget public key, never from the request body."),
    ("2. Scoped data access", "Every repository function requires tenantId; there is no unscoped query helper."),
    ("3. Postgres RLS", "Row-Level Security policies as a safety net if application code ever forgets a filter."),
]
STACK = [
    ("Frontend (dashboard)", "Next.js + TypeScript + Tailwind + shadcn/ui", "Most requested on job listings"),
    ("Widget", "Separate small Preact bundle, one <script> tag", "The embeddable part is what clients pay for"),
    ("Backend", "Node.js + NestJS + TypeScript", "Clean structure, easy to explain in interviews"),
    ("Real-time", "Socket.io with Redis adapter", "Chat, typing indicators, agent presence"),
    ("Database", "PostgreSQL + pgvector (Drizzle ORM)", "One DB for relational data and embeddings"),
    ("Cache / queues", "Redis + BullMQ", "Document ingestion jobs"),
    ("AI", "LLM API + embeddings, RAG pipeline", "Answers come from the client's own data"),
    ("Payments", "Stripe subscriptions", "Shows SaaS billing"),
    ("DevOps", "Docker, GitHub Actions, Vercel + Railway", "Recruiters like seeing CI/CD"),
]
TABLES = [
    ("tenants", "id, name, plan, public_key, branding settings", "One row per client business"),
    ("users", "id, tenant_id, email, password_hash, role (owner / admin / agent)", "Dashboard users and agents"),
    ("documents", "id, tenant_id, title, source, status", "Uploaded knowledge base items"),
    ("document_chunks", "id, document_id, tenant_id, content, embedding (vector)", "Chunks searched by RAG"),
    ("conversations", "id, tenant_id, customer_ref, status (ai / waiting / human / closed), assigned_agent_id", "Chat sessions and handoff state"),
    ("messages", "id, conversation_id, sender_type, content, sources (jsonb), confidence", "Message history with citations"),
    ("subscriptions", "id, tenant_id, stripe_ids, plan, status", "Stripe billing state"),
    ("usage_events", "tenant_id, type, created_at", "Plan limits and analytics"),
]
WORKFLOWS = [
    ("CI (ci.yml)", "Every PR and push to main", "Docs sync check, lint, typecheck, unit tests, integration tests on real Postgres + Redis, build, widget size budget, Docker build, Playwright E2E"),
    ("Deploy (deploy.yml)", "CI success on main", "Migrate DB, deploy API + worker to Railway, dashboard + widget to Vercel, smoke test /health"),
    ("RAG Eval (eval.yml)", "Weekly, manual, PR label run-eval", "Scores retrieval hit@k, answer correctness, correct refusals and injection resistance; fails below the minimum score"),
]
LIFECYCLE = [
    "Pick the next task whose dependencies are done.",
    "Mark it started: python scripts/progress.py start T0X.",
    "Create a branch feat/T0X-short-name and build with tests in small Conventional Commits.",
    "Check every acceptance criterion; run lint, typecheck and tests.",
    "Run the after-every-task checklist (below), then open a PR with the template.",
    "CI must be green; squash merge. From T16 every merge deploys automatically.",
    "At the end of each week tag a release (v0.1.0 ... v1.0.0).",
]
CHECKLIST = [
    "python scripts/progress.py done T0X --log \"summary\" updates tasks.json, PLAN.md, README, the Excel tracker and this Word document.",
    "Expand the entry in docs/PROGRESS.md: files, how to verify, decisions, next.",
    "Update ARCHITECTURE.md if the design changed; PIPELINE.md if CI, env vars or secrets changed.",
    "Re-run python scripts/progress.py sync if those edits should show in this document.",
    "Commit the docs in the same PR as the code.",
]
SECURITY = [
    ("Cross-tenant data access", "Three-layer tenancy + isolation test suite"),
    ("Credential theft", "argon2 hashes, 15-minute access tokens, rotating httpOnly refresh cookie"),
    ("Widget embedded on a foreign site", "Allowed-domains origin check"),
    ("SSRF via URL ingestion", "Block private, loopback and metadata IP ranges; size and time limits"),
    ("Prompt injection", "Input and docs passed as untrusted data, system prompt never echoed, attacks in the eval set"),
    ("PII leakage", "Redaction before LLM calls and in logs"),
    ("Abuse and cost attacks", "Per-IP and per-tenant rate limits, plan limits"),
    ("Forged billing events", "Stripe signature verification, idempotent webhook handling"),
]


# ----------------------------------------------------------------- helpers ---
def shade(cell, hex_fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def cell_margins(table, top=60, bottom=60, left=100, right=100) -> None:
    tbl_pr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for side, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tbl_pr.append(mar)


def borders(table, color="D5D7E3") -> None:
    tbl_pr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), color)
        b.append(el)
    tbl_pr.append(b)


def add_runs(par, text: str, size=None, color=None, bold=False) -> None:
    """Minimal inline markdown: **bold** and `code`."""
    for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
        if not part:
            continue
        if part.startswith("**"):
            r = par.add_run(part[2:-2]); r.bold = True
        elif part.startswith("`"):
            r = par.add_run(part[1:-1]); r.font.name = "Consolas"; r.font.color.rgb = RGBColor(0x3B, 0x3F, 0x8C)
        else:
            r = par.add_run(part); r.bold = bold or None
        if size:
            r.font.size = Pt(size)
        if color is not None and not part.startswith("`"):
            r.font.color.rgb = color


def para(doc, text="", style=None, size=None, color=None, bold=False, space_after=4, align=None):
    p = doc.add_paragraph(style=style)
    add_runs(p, text, size=size, color=color, bold=bold)
    p.paragraph_format.space_after = Pt(space_after)
    if align:
        p.alignment = align
    return p


def bullet(doc, text: str, label: str | None = None):
    p = doc.add_paragraph(style="List Bullet")
    if label:
        r = p.add_run(label + ": "); r.bold = True
    add_runs(p, text)
    p.paragraph_format.space_after = Pt(2)
    return p


def numbered(doc, items):
    for i, text in enumerate(items, 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Mm(8)
        p.paragraph_format.first_line_indent = Mm(-6)
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(f"{i}.  "); r.bold = True; r.font.color.rgb = BRAND
        add_runs(p, text)


def table(doc, headers, rows, widths_mm, status_col=None, font_size=9):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    borders(t)
    cell_margins(t)
    for i, w in enumerate(widths_mm):  # grid widths (LibreOffice / Google Docs use these)
        t.columns[i].width = Mm(w)
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.width = Mm(widths_mm[i])
        shade(c, BRAND_HEX)
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = c.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(h); r.bold = True; r.font.size = Pt(font_size); r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    tr_pr = t.rows[0]._tr.get_or_add_trPr()
    hdr = OxmlElement("w:tblHeader"); hdr.set(qn("w:val"), "true"); tr_pr.append(hdr)
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for i, val in enumerate(row):
            c = cells[i]
            c.width = Mm(widths_mm[i])
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            if status_col is not None and i == status_col:
                shade(c, STATUS_FILL.get(val, "FFFFFF"))
                add_runs(p, str(val), size=font_size, color=STATUS_INK.get(val, INK), bold=True)
            else:
                add_runs(p, str(val), size=font_size, color=INK)
                if ri % 2 == 1:
                    shade(c, "F7F8FB")
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def progress_bar(doc, pct: int, cells=25):
    t = doc.add_table(rows=1, cols=cells)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    filled = round(cells * pct / 100)
    for i, c in enumerate(t.rows[0].cells):
        c.width = Mm(CONTENT_W / cells)
        shade(c, BRAND_HEX if i < filled else "E4E5EE")
        c.paragraphs[0].paragraph_format.space_after = Pt(0)
        c.paragraphs[0].add_run(" ").font.size = Pt(4)
    tr_pr = t.rows[0]._tr.get_or_add_trPr()
    h = OxmlElement("w:trHeight"); h.set(qn("w:val"), "160"); h.set(qn("w:hRule"), "exact"); tr_pr.append(h)
    borders(t, color="FFFFFF")


def page_number_footer(section, label: str):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(label + "  ·  Page "); r.font.size = Pt(8); r.font.color.rgb = MUTED
    run = p.add_run(); run.font.size = Pt(8); run.font.color.rgb = MUTED
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = text
        run._r.append(el)


def heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    return h


def parse_progress(path: Path, limit=6):
    if not path.exists():
        return []
    text = path.read_text(encoding="utf-8").split("<!-- NEW ENTRIES BELOW -->", 1)[-1]
    entries = []
    for block in re.split(r"^### ", text, flags=re.M)[1:]:
        lines = block.strip().splitlines()
        title = lines[0].strip()
        items = [ln.strip()[2:] for ln in lines[1:] if ln.strip().startswith("- ")]
        entries.append((title, items))
    return entries[:limit]


# ------------------------------------------------------------------ build ---
def build(data: dict, progress_md: Path, out: Path) -> None:
    tasks = data["tasks"]
    done = sum(t["status"] == "Done" for t in tasks)
    wip = [t for t in tasks if t["status"] == "In progress"]
    pct = round(100 * done / len(tasks))
    nxt = wip[0] if wip else next((t for t in tasks if t["status"] == "Not started"), None)
    today = dt.date.today().strftime("%d %b %Y")

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    for side in ("left_margin", "right_margin"):
        setattr(sec, side, Mm(20))
    sec.top_margin, sec.bottom_margin = Mm(20), Mm(18)
    sec.different_first_page_header_footer = True

    st = doc.styles
    st["Normal"].font.name = "Calibri"
    st["Normal"].font.size = Pt(10.5)
    st["Normal"].font.color.rgb = INK
    st["Normal"].element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    for lvl, size in ((1, 17), (2, 13), (3, 11.5)):
        s = st[f"Heading {lvl}"]
        s.font.name = "Calibri"; s.font.size = Pt(size); s.font.bold = True
        s.font.color.rgb = BRAND if lvl < 3 else INK
        rfonts = s.element.rPr.rFonts
        for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            if rfonts.get(qn(attr)) is not None:
                del rfonts.attrib[qn(attr)]
        rfonts.set(qn("w:ascii"), "Calibri"); rfonts.set(qn("w:hAnsi"), "Calibri")
        s.paragraph_format.space_before = Pt(16 if lvl == 1 else 10)
        s.paragraph_format.space_after = Pt(6 if lvl == 1 else 4)

    page_number_footer(sec, "SupportPilot  ·  Project Document")

    # ---- cover
    for _ in range(5):
        doc.add_paragraph()
    p = para(doc, "PROJECT DOCUMENT", size=10, color=MUTED, bold=True, space_after=6)
    p = doc.add_paragraph(); r = p.add_run("SupportPilot"); r.bold = True; r.font.size = Pt(40); r.font.color.rgb = BRAND
    p.paragraph_format.space_after = Pt(4)
    para(doc, data.get("tagline", ""), size=14, color=INK, space_after=24)
    progress_bar(doc, pct)
    para(doc, f"**{done} of {len(tasks)} build tasks done ({pct}%)**" + (f"  ·  current task: {nxt['id']} {nxt['title']}" if nxt else "  ·  v1.0 complete"),
         size=10, color=MUTED, space_after=2)
    para(doc, f"Plan: 6 weeks part time, starting {dt.date.fromisoformat(data['plan_start']).strftime('%d %b %Y')}", size=10, color=MUTED, space_after=2)
    para(doc, f"Prepared by Akash  ·  Updated {today}", size=10, color=MUTED, space_after=2)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # ---- contents
    heading(doc, "Contents", 1)
    for i, name in enumerate(["Executive summary", "Problem and solution", "Features", "How it works",
                              "Architecture", "Tech stack", "Data model", "Build plan and progress",
                              "Task breakdown", "CI/CD pipeline", "Development workflow", "Security",
                              "Recent progress"], 1):
        para(doc, f"{i}.  {name}", space_after=2)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # ---- 1 summary
    heading(doc, "1. Executive summary", 1)
    para(doc, SUMMARY, space_after=8)
    for label, text in AUDIENCE:
        bullet(doc, text, label)

    # ---- 2 problem / solution
    heading(doc, "2. Problem and solution", 1)
    heading(doc, "The problem", 2)
    for t_ in PROBLEM:
        bullet(doc, t_)
    heading(doc, "The solution", 2)
    for t_ in SOLUTION:
        bullet(doc, t_)

    # ---- 3 features
    heading(doc, "3. Features", 1)
    table(doc, ["Feature", "What it does"], FEATURES, [45, CONTENT_W - 45])
    heading(doc, "Standout features", 2)
    para(doc, "What separates SupportPilot from a tutorial chatbot:", space_after=4)
    table(doc, ["Feature", "Why it stands out", "Task"], STANDOUT, [52, CONTENT_W - 52 - 22, 22])

    # ---- 4 how it works
    heading(doc, "4. How it works", 1)
    heading(doc, "Chat flow", 2)
    numbered(doc, CHAT_FLOW)
    heading(doc, "Conversation states (human handoff)", 2)
    table(doc, ["State", "Meaning", "Can move to"], STATES, [22, 78, CONTENT_W - 100])
    para(doc, "Confidence combines retrieval similarity of the cited chunks with the model's own rating. "
              "The threshold is set per tenant (default 0.55); two low-confidence answers in a row also hand off.",
         size=9.5, color=MUTED)

    # ---- 5 architecture
    heading(doc, "5. Architecture", 1)
    table(doc, ["Component", "Responsibility", "Runs on"], COMPONENTS, [38, 88, CONTENT_W - 126])
    heading(doc, "Multi-tenancy", 2)
    para(doc, "Rule: every query is scoped by tenant_id. This is enforced in three layers:", space_after=4)
    table(doc, ["Layer", "How"], TENANCY, [42, CONTENT_W - 42])
    para(doc, "Full diagrams (system, sequence, ingestion, state machine, ER) are in docs/ARCHITECTURE.md in the repository.",
         size=9.5, color=MUTED)

    # ---- 6 stack
    heading(doc, "6. Tech stack", 1)
    table(doc, ["Layer", "Choice", "Why"], STACK, [36, 70, CONTENT_W - 106])

    # ---- 7 data model
    heading(doc, "7. Data model", 1)
    table(doc, ["Table", "Key columns", "Purpose"], TABLES, [34, 88, CONTENT_W - 122])

    # ---- 8 plan & progress
    heading(doc, "8. Build plan and progress", 1)
    progress_bar(doc, pct)
    para(doc, f"**{done} / {len(tasks)} tasks done ({pct}%)**"
              + (f"  ·  current / next: {nxt['id']} {nxt['title']} ({nxt['status']})" if nxt else ""),
         size=10, space_after=8)
    heading(doc, "Milestones", 2)
    rows = []
    for w in data["weeks"]:
        wt = [t for t in tasks if t["week"] == w["week"]]
        d = sum(t["status"] == "Done" for t in wt)
        rows.append((w["week"], w["dates"], f"{w['milestone']} ({w['version']})", w["outcome"], f"{d}/{len(wt)}"))
    table(doc, ["Week", "Dates", "Milestone", "Outcome", "Done"], rows, [16, 28, 36, CONTENT_W - 96, 16])
    heading(doc, "Task status", 2)
    table(doc, ["ID", "Week", "Task", "Est.", "Status", "Completed"],
          [(t["id"], t["week"], t["title"], f"{t.get('est_hours', '')} h", t["status"], t.get("completed") or "")
           for t in tasks],
          [13, 18, 70, 13, 28, CONTENT_W - 142], status_col=4)

    # ---- 9 task breakdown
    heading(doc, "9. Task breakdown", 1)
    para(doc, "Every task with its goal, subtasks and acceptance criteria. Checked items are complete.",
         size=9.5, color=MUTED)
    for w in data["weeks"]:
        heading(doc, f"{w['week']} ({w['dates']}): {w['milestone']}", 2)
        for t in [t for t in tasks if t["week"] == w["week"]]:
            h = heading(doc, f"{t['id']}  {t['title']}", 3)
            meta = doc.add_paragraph()
            r = meta.add_run(f" {t['status']} "); r.bold = True; r.font.size = Pt(8.5)
            r.font.color.rgb = STATUS_INK[t["status"]]
            rpr = r._r.get_or_add_rPr(); shd = OxmlElement("w:shd")
            shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), STATUS_FILL[t["status"]]); rpr.append(shd)
            deps = ", ".join(t.get("depends") or []) or "none"
            r2 = meta.add_run(f"   est. {t.get('est_hours')} h  ·  depends on: {deps}"
                              + (f"  ·  completed {t['completed']}" if t.get("completed") else ""))
            r2.font.size = Pt(8.5); r2.font.color.rgb = MUTED
            meta.paragraph_format.space_after = Pt(4)
            para(doc, f"**Goal:** {t['goal']}", space_after=3)
            is_done = t["status"] == "Done"
            para(doc, "**Subtasks**", space_after=1)
            for s in t["subtasks"]:
                ticked = is_done or s.startswith("[x] ")
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Mm(6); p.paragraph_format.first_line_indent = Mm(-5)
                p.paragraph_format.space_after = Pt(1)
                box = p.add_run("☑ " if ticked else "☐ "); box.font.color.rgb = STATUS_INK["Done"] if ticked else MUTED
                add_runs(p, s.removeprefix("[x] "), size=9.5)
            para(doc, "**Acceptance criteria**", space_after=1)
            for a in t["acceptance"]:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Mm(6); p.paragraph_format.first_line_indent = Mm(-5)
                p.paragraph_format.space_after = Pt(1)
                box = p.add_run("☑ " if is_done else "☐ "); box.font.color.rgb = STATUS_INK["Done"] if is_done else MUTED
                add_runs(p, a, size=9.5)
            extra = f"**Deliverables:** {', '.join(t['deliverables'])}"
            para(doc, extra, size=9.5, space_after=1)
            if t.get("standout"):
                para(doc, f"**Standout feature:** {t['standout']}", size=9.5, space_after=1)
            para(doc, f"**Interview talking point:** {t['talking_point']}", size=9.5, color=MUTED, space_after=8)

    # ---- 10 pipeline
    heading(doc, "10. CI/CD pipeline", 1)
    para(doc, "Three GitHub Actions workflows. Jobs that need application code switch on automatically as the code lands, "
              "so the pipeline is green from the first commit.", space_after=6)
    table(doc, ["Workflow", "Trigger", "What it does"], WORKFLOWS, [36, 40, CONTENT_W - 76])
    heading(doc, "Stages on every pull request", 2)
    numbered(doc, [
        "Docs check: PLAN.md and README must match tasks.json (progress sync was run).",
        "Quality: lint, typecheck and unit tests with coverage.",
        "Integration: API tests against real PostgreSQL (pgvector) and Redis containers; AI providers mocked.",
        "Build: all apps; widget.js must stay under a 60 KB gzipped budget.",
        "Docker: API and worker images build.",
        "End-to-end: Playwright runs signup -> upload -> widget chat -> handoff -> agent reply.",
    ])
    heading(doc, "Deploy (after merge to main)", 2)
    numbered(doc, [
        "Run database migrations (backwards compatible: expand, deploy, contract).",
        "Deploy API and worker to Railway.",
        "Deploy dashboard and widget to Vercel.",
        "Smoke test /health and the dashboard URL. Rollback: Vercel instant rollback, Railway redeploy previous; database is fix-forward.",
    ])

    # ---- 11 workflow
    heading(doc, "11. Development workflow", 1)
    heading(doc, "Task lifecycle", 2)
    numbered(doc, LIFECYCLE)
    heading(doc, "After every completed task", 2)
    numbered(doc, CHECKLIST)
    para(doc, "Definition of Done: acceptance criteria met, tests cover the new logic, every query tenant-scoped, CI green, "
              "no secrets in code, docs updated.", size=9.5, color=MUTED)

    # ---- 12 security
    heading(doc, "12. Security", 1)
    table(doc, ["Threat", "Control"], SECURITY, [55, CONTENT_W - 55])

    # ---- 13 progress log
    heading(doc, "13. Recent progress", 1)
    entries = parse_progress(progress_md)
    if not entries:
        para(doc, "No entries yet.", color=MUTED)
    for title, items in entries:
        p = para(doc, f"**{title}**", space_after=2)
        for it in items:
            m = re.match(r"\*\*(.+?):\*\*\s*(.*)", it)
            if m:
                bullet(doc, m.group(2), m.group(1))
            else:
                bullet(doc, it)
        doc.add_paragraph().paragraph_format.space_after = Pt(2)

    cp = doc.core_properties
    cp.title = "SupportPilot - Project Document"
    cp.author = "Akash"
    cp.subject = data.get("tagline", "")
    cp.keywords = "SupportPilot, RAG, SaaS, multi-tenant, NestJS, Next.js, pgvector"
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)


if __name__ == "__main__":
    data = json.loads((ROOT / "docs" / "plan" / "tasks.json").read_text(encoding="utf-8"))
    out = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT.parent / "SupportPilot_Project_Document.docx"
    build(data, ROOT / "docs" / "PROGRESS.md", out)
    print(f"wrote {out}")
