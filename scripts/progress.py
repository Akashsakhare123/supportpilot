#!/usr/bin/env python3
"""SupportPilot progress tracker.

Single source of truth: docs/plan/tasks.json
After every task, run one command and every progress view is regenerated:

    python scripts/progress.py start T04          # mark In progress (+ sync)
    python scripts/progress.py done T04 --log "Upload API + KB page"   # mark Done (+ sync)
    python scripts/progress.py sync               # regenerate everything
    python scripts/progress.py status             # print a summary

`sync` updates:
  - docs/PLAN.md            (AUTO:STATUS, AUTO:MILESTONES, AUTO:TASKS blocks)
  - README.md               (AUTO:PROGRESS block)
  - ../SupportPilot_Build_Plan.xlsx   (Status + Completed on columns)
  - ../SupportPilot_Project_Document.docx  (regenerated Word document)

Requires: Python 3.9+, openpyxl, python-docx  (pip install openpyxl python-docx)
"""
from __future__ import annotations

import argparse
import datetime as dt
import json
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
TASKS_JSON = ROOT / "docs" / "plan" / "tasks.json"
PLAN_MD = ROOT / "docs" / "PLAN.md"
README_MD = ROOT / "README.md"
PROGRESS_MD = ROOT / "docs" / "PROGRESS.md"
DEFAULT_XLSX = ROOT.parent / "SupportPilot_Build_Plan.xlsx"
DEFAULT_DOCX = ROOT.parent / "SupportPilot_Project_Document.docx"

STATUSES = ("Not started", "In progress", "Done")


# ---------------------------------------------------------------- data ----
def load() -> dict:
    return json.loads(TASKS_JSON.read_text(encoding="utf-8"))


def save(data: dict) -> None:
    TASKS_JSON.write_text(json.dumps(data, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")


def find(data: dict, task_id: str) -> dict:
    task_id = task_id.upper()
    for t in data["tasks"]:
        if t["id"] == task_id:
            return t
    sys.exit(f"Unknown task id: {task_id}")


def stats(data: dict) -> dict:
    tasks = data["tasks"]
    done = [t for t in tasks if t["status"] == "Done"]
    wip = [t for t in tasks if t["status"] == "In progress"]
    total_h = sum(t.get("est_hours", 0) for t in tasks)
    done_h = sum(t.get("est_hours", 0) for t in done)
    nxt = wip[0] if wip else next((t for t in tasks if t["status"] == "Not started"), None)
    return {
        "total": len(tasks), "done": len(done), "wip": len(wip),
        "pct": round(100 * len(done) / len(tasks)) if tasks else 0,
        "total_h": total_h, "done_h": done_h, "next": nxt,
    }


def last_update(data: dict) -> str:
    """Latest start/complete date in tasks.json (deterministic, so CI can diff the docs)."""
    dates = [d for t in data["tasks"] for d in (t.get("started"), t.get("completed")) if d]
    return max(dates) if dates else "not started yet"


def bar(pct: int, width: int = 20) -> str:
    filled = round(width * pct / 100)
    return "[" + "#" * filled + "-" * (width - filled) + f"] {pct}%"


def week_progress(data: dict, week: str) -> tuple[int, int]:
    ts = [t for t in data["tasks"] if t["week"] == week]
    return sum(t["status"] == "Done" for t in ts), len(ts)


# -------------------------------------------------------------- render ----
def render_status(data: dict) -> str:
    s = stats(data)
    nxt = s["next"]
    lines = [
        f"`{bar(s['pct'])}`  **{s['done']} / {s['total']} tasks done** "
        f"({s['done_h']} / {s['total_h']} est. hours)  ",
        f"Last update: {last_update(data)}  ",
        f"**Current / next task:** {nxt['id']} {nxt['title']} ({nxt['status']})" if nxt else "**All tasks done.**",
        "",
        "| ID | Week | Task | Status | Completed |",
        "|---|---|---|---|---|",
    ]
    for t in data["tasks"]:
        status = f"**{t['status']}**" if t["status"] == "In progress" else t["status"]
        lines.append(f"| [{t['id']}](#{anchor(t)}) | {t['week']} | {t['title']} | {status} | {t.get('completed') or ''} |")
    return "\n".join(lines)


def render_milestones(data: dict) -> str:
    lines = ["| Week | Dates | Milestone | Release | Outcome | Progress |", "|---|---|---|---|---|---|"]
    for w in data["weeks"]:
        d, n = week_progress(data, w["week"])
        lines.append(f"| {w['week']} | {w['dates']} | {w['milestone']} | {w['version']} | {w['outcome']} | {d}/{n} |")
    return "\n".join(lines)


def anchor(t: dict) -> str:
    text = f"{t['id']} - {t['title']}".lower()
    text = re.sub(r"[^a-z0-9 \-]", "", text)
    return text.replace(" ", "-")


def render_tasks(data: dict) -> str:
    out = []
    for w in data["weeks"]:
        out.append(f"### {w['week']} ({w['dates']}): {w['milestone']}")
        out.append("")
        for t in [t for t in data["tasks"] if t["week"] == w["week"]]:
            done = t["status"] == "Done"
            deps = ", ".join(t.get("depends") or []) or "none"
            out += [
                f"#### {t['id']} - {t['title']}",
                "",
                f"`{t['status']}` · est. {t.get('est_hours', '?')} h · depends on: {deps}"
                + (f" · completed {t['completed']}" if t.get("completed") else ""),
                "",
                f"**Goal:** {t['goal']}",
                "",
                "**Subtasks**",
                "",
            ]
            for st in t["subtasks"]:
                ticked = done or st.startswith("[x] ")
                out.append(f"- [{'x' if ticked else ' '}] {st.removeprefix('[x] ')}")
            out += ["", "**Acceptance criteria**", ""]
            for a in t["acceptance"]:
                out.append(f"- [{'x' if done else ' '}] {a}")
            out += ["", f"**Deliverables:** {', '.join(t['deliverables'])}  "]
            if t.get("standout"):
                out.append(f"**Standout feature:** {t['standout']}  ")
            out += [f"**Interview talking point:** {t['talking_point']}", ""]
    return "\n".join(out).rstrip()


def render_readme(data: dict) -> str:
    s = stats(data)
    lines = [
        f"![progress](https://img.shields.io/badge/progress-{s['done']}%2F{s['total']}%20tasks-3B3F8C)",
        "",
        f"`{bar(s['pct'])}`" + (f" - currently on **{s['next']['id']} {s['next']['title']}**" if s["next"] else " - v1.0 complete"),
        "",
        "| Milestone | Release | Progress |",
        "|---|---|---|",
    ]
    for w in data["weeks"]:
        d, n = week_progress(data, w["week"])
        mark = "done" if d == n else ("in progress" if d or any(
            t["status"] == "In progress" for t in data["tasks"] if t["week"] == w["week"]) else "planned")
        lines.append(f"| {w['milestone']} | {w['version']} | {d}/{n} ({mark}) |")
    lines += ["", "Full plan: [docs/PLAN.md](docs/PLAN.md) · Build log: [docs/PROGRESS.md](docs/PROGRESS.md)"]
    return "\n".join(lines)


def replace_block(text: str, name: str, body: str) -> str:
    pattern = re.compile(rf"(<!-- AUTO:{name} -->)(.*?)(<!-- /AUTO:{name} -->)", re.S)
    if not pattern.search(text):
        raise SystemExit(f"Marker AUTO:{name} not found")
    return pattern.sub(lambda m: f"{m.group(1)}\n{body}\n{m.group(3)}", text)


# ---------------------------------------------------------------- xlsx ----
def sync_xlsx(data: dict, path: Path) -> None:
    try:
        import openpyxl
        from copy import copy
        from openpyxl.formatting.rule import CellIsRule
        from openpyxl.styles import PatternFill, Font
    except ImportError:
        print("  ! openpyxl not installed, skipping xlsx (pip install openpyxl)")
        return
    if not path.exists():
        print(f"  ! {path.name} not found, skipping xlsx")
        return
    wb = openpyxl.load_workbook(path)
    ws = wb["Build Plan"]
    by_title = {t["title"].strip().lower(): t for t in data["tasks"]}
    by_id = {t["id"]: t for t in data["tasks"]}

    # One-time setup: ID + Completed on columns, status colours.
    if ws["E1"].value != "ID":
        for col, label in (("E", "ID"), ("F", "Completed on")):
            cell = ws[f"{col}1"]
            cell.value = label
            src = ws["D1"]
            cell.font, cell.fill, cell.alignment, cell.border = copy(src.font), copy(src.fill), copy(src.alignment), copy(src.border)
        ws.conditional_formatting.add("D2:D18", CellIsRule(operator="equal", formula=['"Done"'],
                                      fill=PatternFill("solid", fgColor="D9F2E3"), font=Font(color="1E6B3A", bold=True)))
        ws.conditional_formatting.add("D2:D18", CellIsRule(operator="equal", formula=['"In progress"'],
                                      fill=PatternFill("solid", fgColor="FFF1CC"), font=Font(color="8A5A00", bold=True)))
    updated = 0
    for row in range(2, 19):
        title = (ws[f"B{row}"].value or "").strip().lower()
        t = by_id.get(ws[f"E{row}"].value) or by_title.get(title)
        if not t:
            continue
        ws[f"E{row}"].value = t["id"]
        ws[f"D{row}"].value = t["status"]
        ws[f"F{row}"].value = t.get("completed") or None
        for col in "EF":
            ws[f"{col}{row}"].font = copy(ws[f"D{row}"].font)
            ws[f"{col}{row}"].border = copy(ws[f"D{row}"].border)
            ws[f"{col}{row}"].alignment = copy(ws[f"D{row}"].alignment)
        updated += 1
    wb.calculation.fullCalcOnLoad = True
    try:
        wb.save(path)
        print(f"  - {path.name}: {updated} task rows synced")
    except PermissionError:
        print(f"  ! {path.name} is open in Excel; close it and run sync again")


# ---------------------------------------------------------------- main ----
def sync(args) -> None:
    data = load()
    plan = PLAN_MD.read_text(encoding="utf-8")
    plan = replace_block(plan, "STATUS", render_status(data))
    plan = replace_block(plan, "MILESTONES", render_milestones(data))
    plan = replace_block(plan, "TASKS", render_tasks(data))
    PLAN_MD.write_text(plan, encoding="utf-8")
    print("  - docs/PLAN.md regenerated")
    if README_MD.exists():
        README_MD.write_text(replace_block(README_MD.read_text(encoding="utf-8"), "PROGRESS", render_readme(data)), encoding="utf-8")
        print("  - README.md progress block regenerated")
    if not args.skip_xlsx:
        sync_xlsx(data, Path(args.xlsx))
    if not args.skip_docx:
        sys.dont_write_bytecode = True
        sys.path.insert(0, str(Path(__file__).resolve().parent))
        try:
            import build_docx
        except ImportError as e:
            print(f"  ! Word doc skipped ({e}); pip install python-docx")
        else:
            build_docx.build(data, PROGRESS_MD, Path(args.docx))
            print(f"  - {Path(args.docx).name} regenerated")
    s = stats(data)
    print(f"Progress: {bar(s['pct'])}  {s['done']}/{s['total']} tasks")


def add_log(task: dict, summary: str) -> None:
    text = PROGRESS_MD.read_text(encoding="utf-8")
    entry = (f"### {dt.date.today().isoformat()} · {task['id']} {task['title']}\n\n"
             f"- **Built:** {summary}\n- **Verify:** see acceptance criteria in PLAN.md\n\n")
    marker = "<!-- NEW ENTRIES BELOW -->"
    text = text.replace(marker, marker + "\n\n" + entry.rstrip() + "\n", 1)
    PROGRESS_MD.write_text(text, encoding="utf-8")


def main() -> None:
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = p.add_subparsers(dest="cmd", required=True)
    for name in ("start", "done", "reset"):
        sp = sub.add_parser(name)
        sp.add_argument("task_id")
        sp.add_argument("--date", default=dt.date.today().isoformat())
        sp.add_argument("--log", help="one-line summary appended to docs/PROGRESS.md (done only)")
        sp.add_argument("--no-sync", action="store_true")
    sub.add_parser("status")
    for sp in [sub.choices[n] for n in ("start", "done", "reset")] + [sub.add_parser("sync")]:
        sp.add_argument("--xlsx", default=str(DEFAULT_XLSX))
        sp.add_argument("--docx", default=str(DEFAULT_DOCX))
        sp.add_argument("--skip-xlsx", action="store_true")
        sp.add_argument("--skip-docx", action="store_true")
    args = p.parse_args()

    if args.cmd == "status":
        data = load()
        s = stats(data)
        print(f"{bar(s['pct'])}  {s['done']}/{s['total']} done, {s['wip']} in progress")
        for t in data["tasks"]:
            print(f"  {t['id']}  {t['status']:<12} {t['week']}  {t['title']}")
        return
    if args.cmd == "sync":
        sync(args)
        return

    data = load()
    t = find(data, args.task_id)
    if args.cmd == "start":
        t["status"], t["started"] = "In progress", t.get("started") or args.date
    elif args.cmd == "done":
        t["status"], t["completed"] = "Done", args.date
        t["started"] = t.get("started") or args.date
        if args.log:
            add_log(t, args.log)
    else:
        t["status"], t["started"], t["completed"] = "Not started", None, None
    save(data)
    print(f"{t['id']} -> {t['status']}")
    if not args.no_sync:
        sync(args)


if __name__ == "__main__":
    main()
